"""Realistic fictional proposal/report, with reusable native document scenes."""
import json
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from docx import Document
from docx.shared import Mm, Pt, RGBColor
import fitz
from scenes import Scene


def base(b,kind,num,title,subtitle='',dark=False):
    s=Scene(b,595.276,841.89,'ink' if dark else 'paper')
    s.logo(42,35,140,'white' if dark else 'color')
    s.text(('PROPUESTA' if kind=='proposal' else 'INFORME')+' / EJEMPLO',374,43,179,8,'petal' if dark else 'muted','Medium')
    s.line(42,81,553,81,'muted' if dark else 'line')
    if title:s.text(title,42,116,510,30,'white' if dark else 'ink','Medium',1.16)
    if subtitle:s.text(subtitle,42,207,499,11,'petal' if dark else 'muted')
    s.line(42,779,553,779,'muted' if dark else 'line')
    s.text('TENBELTZ / LADERA CLOUD / EJEMPLO FICTICIO',42,794,450,7.5,'petal' if dark else 'muted')
    s.text(f'{num:02d}',527,794,25,8,'petal' if dark else 'muted')
    return s


def block(s,title,body,x,y,w=235):
    s.text(title,x,y,w,15,'accent','Medium')
    return s.text(body,x,y+32,w,11,'muted',leading=1.5)


def note(s,text,y=719):
    s.box(42,y,511,43,'surface');s.text(text,54,y+11,488,9,'muted',leading=1.35)


def table(s,headers,rows,x,y,widths,height=61):
    for i,row in enumerate([headers]+rows):
        xx=x;yy=y+(34 if i else 0)+(i-1)*height if i else y
        for j,text in enumerate(row):
            w=widths[j];s.box(xx,yy,w-2,32 if i==0 else height-2,'accent' if i==0 else ('surface' if i%2 else 'paper'))
            s.text(str(text),xx+10,yy+10,w-20,9 if i==0 else 10,'white' if i==0 else 'ink','Medium' if i==0 else 'Regular',1.4)
            xx+=w


def architecture(s,y=290):
    nodes=[('Producto SaaS','Widget + sesión'),('API del asistente','Permisos + límites'),('Recuperación','Documentos + citas')]
    for j,(title,sub) in enumerate(nodes):
        x=42+j*178;s.box(x,y,155,94,'surface');s.text(title,x+12,y+17,130,12,'ink','Medium');s.text(sub,x+12,y+51,130,9,'muted')
        if j<2:s.arrow(x+159,y+47,x+171,y+47,stroke=1)
    s.arrow(475,y+94,475,y+133,stroke=1)
    s.box(398,y+140,155,94,'accent');s.text('Modelo de IA',410,y+158,130,12,'white','Medium');s.text('Respuesta + evidencia',410,y+191,130,9,'white')
    s.box(42,y+140,323,94,'surface');s.text('Base documental del cliente',55,y+157,290,12,'ink','Medium');s.text('Versionado, permisos y contenido validado.',55,y+188,290,10,'muted')
    s.arrow(365,y+187,393,y+187,stroke=1)
    s.box(42,y+272,511,57,'ink');s.text('Evaluación + trazas + supervisión humana',56,y+291,480,12,'white','Medium')
    s.text('Flujo simplificado / arquitectura conceptual del ejemplo',42,y+345,510,8,'muted')


def proposal(b,d):
    p=d['proposal'];sc=[]
    s=base(b,'proposal',1,'','')
    s.text('PROPUESTA TÉCNICA / 2026.10',42,137,510,9,'accent','Medium')
    s.text('Un asistente\nde soporte,\nintegrado en\ntu SaaS.',42,191,420,43,'ink','Medium',1.1)
    s.image(b.DIST/'graphics'/'tenbeltz-flower-transparent.png',286,372,293,293)
    s.text('Ladera Cloud',42,478,280,21,'accent','Medium')
    s.text('Piloto de seis semanas para responder con documentación, citas y un alcance evaluable.',42,525,258,12,'muted',leading=1.5)
    s.text('08 octubre 2026 / Versión 1.0\nResponsable técnico: TenBeltz',42,672,440,10,'muted')
    note(s,'Cliente, presupuesto y alcance ficticios. Ejemplo de diseño; no es una oferta comercial.',726);sc.append(s)
    s=base(b,'proposal',2,'De la documentación\na una respuesta útil.','Contexto del producto y alcance de la primera entrega.')
    s.text('El reto',42,277,490,17,'accent','Medium')
    s.text('Ladera Cloud centraliza operaciones de equipos de servicios. Su soporte responde consultas sobre configuración, permisos y flujos de trabajo con una documentación distribuida. El objetivo del piloto es integrar respuestas trazables dentro del producto.',42,315,499,12,'ink',leading=1.55)
    s.line(42,416,553,416)
    block(s,'Incluido en el piloto','Búsqueda documental, respuestas con citas, control de permisos, abstención cuando no hay evidencia y registro de trazas.',42,455)
    block(s,'Fuera de alcance','Acciones autónomas sobre datos, cambios en producción, atención multilingüe y automatización de decisiones comerciales.',316,455)
    s.box(42,616,511,83,'accent');s.text('Éxito definido antes de construir.',56,631,480,17,'white','Medium');s.text('Acordar muestra, rúbrica, umbrales y límites con el responsable de producto.',56,665,480,10,'white')
    note(s,'Dependencias: documentación aprobada, entorno de prueba y un interlocutor técnico del cliente.');sc.append(s)
    s=base(b,'proposal',3,'Una arquitectura\nque se puede evaluar.','Integración gradual, permisos explícitos y observabilidad desde el piloto.')
    architecture(s,279)
    block(s,'Diseño de la recuperación','El cliente valida las fuentes y sus permisos. El asistente sólo utiliza documentación autorizada para la sesión.',42,654,w=240)
    block(s,'Decisión de lanzamiento','La evaluación y la revisión humana preceden a la apertura a usuarios. La entrega incluye límites documentados.',316,654,w=235)
    sc.append(s)
    s=base(b,'proposal',4,'Seis semanas.\nTres entregas verificables.','Calendario orientativo del ejemplo, sujeto a disponibilidad de datos y revisión.')
    x0=236;cell=50
    for week in range(6):s.text(f'S{week+1}',x0+week*cell+13,286,43,9,'muted','Medium');s.line(x0+week*cell,317,x0+week*cell,500)
    for i,phase in enumerate(p['phases']):
        yy=335+i*61;s.text(phase['name'],42,yy,184,12,'ink','Medium')
        for week in range(i*2,i*2+2):s.box(x0+week*cell+2,yy-2,46,29,'accent' if i!=1 else 'petal')
    for i,phase in enumerate(p['phases']):
        yy=550+i*63;s.text(f'0{i+1}',42,yy,50,13,'accent','Medium');s.text(phase['deliverable'],98,yy,446,10,'muted',leading=1.4)
    note(s,'Cada fase termina con una revisión conjunta. El calendario se ajusta si cambian fuentes o alcance.');sc.append(s)
    s=base(b,'proposal',5,'Una inversión\nligada a entregables.','Presupuesto ilustrativo para mostrar cómo presentar una propuesta.')
    s.text('12.000 €',42,276,495,57,'accent','Medium');s.text('Importe de ejemplo / impuestos no incluidos',42,354,490,10,'muted')
    table(s,['Fase','Entregable','Importe'],[
      ['Definición','Datos, rúbrica y arquitectura','2.400 €'],['Construcción','Piloto integrado y citas','6.000 €'],['Entrega','Evaluación y operación','3.600 €']],42,410,[153,254,104],67)
    block(s,'Responsabilidades','TenBeltz aporta dirección técnica e implementación. Ladera aporta fuentes, permisos, revisión de contenido y acceso al entorno de prueba.',42,663,w=511)
    sc.append(s)
    s=base(b,'proposal',6,'Entregar también\nes explicar los límites.','Criterios de aceptación y siguiente paso.')
    table(s,['Criterio','Cómo se verifica'],[
      ['Calidad','Muestra acordada de 120 casos; objetivo de ejemplo ≥90 % según rúbrica.'],
      ['Latencia','p95 por debajo de 5 s en el entorno y carga acordados.'],
      ['Control','Citas, permisos, abstención y revisión de casos fuera de alcance.'],
      ['Operación','Trazas, documentación y sesión de transferencia al equipo.']],42,286,[135,376],74)
    s.text('Primero, revisar el contexto juntos.',42,651,495,20,'accent','Medium');s.text('Confirmar fuentes, responsables y criterios antes de cerrar el alcance.\nhello@tenbeltz.com / tenbeltz.com',42,696,495,11,'muted',leading=1.55)
    sc.append(s)
    return sc


def horizontal_bar(s,label,value,total,x,y,w=511,fill='accent',detail=None):
    s.text(label,x,y,w,11,'ink','Medium');s.box(x,y+26,w,18,'surface');s.box(x,y+26,w*value/total,18,fill)
    s.text(detail or str(value),x,y+54,w,9,'muted')


def report(b,d):
    e=d['evaluation'];pass_pct=100*e['pilotPass']/e['sampleSize'];baseline_pct=100*e['baselinePass']/e['sampleSize'];sc=[]
    s=base(b,'report',1,'','')
    s.text('INFORME DE EVALUACIÓN / 2026.10',42,137,510,9,'accent','Medium')
    s.text('Del piloto\na la decisión\nde entrega.',42,204,490,45,'ink','Medium',1.12)
    s.image(b.DIST/'graphics'/'tenbeltz-background-petal-contours-paper.png',172,394,400,225)
    s.text('Ladera Cloud',42,588,470,22,'accent','Medium');s.text('Asistente de soporte con respuestas verificables.\nEvaluación de una muestra ficticia de 120 casos.',42,638,490,12,'muted',leading=1.5)
    note(s,'Todas las métricas son sintéticas y sólo ilustran el diseño del informe. No son resultados de TenBeltz.',726);sc.append(s)
    s=base(b,'report',2,'La calidad mejora.\nEl piloto necesita una iteración.','Resumen ejecutivo / datos sintéticos del ejemplo.')
    metrics=[(f'{pass_pct:.1f} %'.replace('.',','),'Casos que pasan la rúbrica'),('120','Casos evaluados'),('4,6 s','Latencia p95'),('14','Casos por resolver')]
    for i,(value,label) in enumerate(metrics):
        x=42+(i%2)*270;y=276+(i//2)*144;s.box(x,y,241,118,'surface');s.text(value,x+16,y+17,210,34,'accent','Medium');s.text(label,x+16,y+77,210,10,'muted')
    s.box(42,592,511,98,'accent');s.text('Decisión del ejemplo',56,607,481,11,'white','Medium');s.text('Continuar con un piloto controlado.',56,633,481,20,'white','Medium');s.text('El objetivo de calidad del 90 % todavía no se alcanza.',56,671,480,10,'white')
    note(s,'Mejora de 18,3 puntos porcentuales frente a la base. Ambas versiones usan los mismos 120 casos.');sc.append(s)
    s=base(b,'report',3,'Evaluar el uso real,\ncon una muestra explícita.','Método, alcance y trazabilidad de las conclusiones.')
    table(s,['Segmento','Casos','Qué se comprueba'],[
      ['Consultas frecuentes','60','Respuestas comunes y documentación conocida.'],
      ['Casos complejos','40','Ambigüedad, composición de pasos y permisos.'],
      ['Fuera de alcance','20','Abstención y derivación al soporte humano.']],42,279,[176,62,273],78)
    block(s,'Rúbrica conjunta','Cada caso pasa sólo si la respuesta es correcta, está sustentada y respeta el alcance. La evaluación del ejemplo usa revisión humana y registro por caso.',42,595,w=239)
    block(s,'Límites de la muestra','No representa tráfico real ni acredita generalización. Los valores son sintéticos. No se estiman intervalos de confianza con estos datos de diseño.',315,595,w=238)
    sc.append(s)
    s=base(b,'report',4,'Más respuestas útiles.\nLa misma muestra.', 'Calidad global y por segmento / comparación emparejada del ejemplo.')
    horizontal_bar(s,'Versión base',baseline_pct,100,42,283,detail='84 de 120 casos / 70,0 %',fill='muted')
    horizontal_bar(s,'Piloto con recuperación documental',pass_pct,100,42,379,detail='106 de 120 casos / 88,3 %')
    s.text('Calidad por segmento',42,493,510,17,'ink','Medium')
    for i,c in enumerate(e['categories']):
        yy=540+i*58;s.text(c['name'],42,yy,188,10,'ink')
        xx=244;ww=236;s.box(xx,yy+1,ww,12,'surface');s.box(xx,yy+1,ww*c['pilotPass']/c['total'],12,'accent')
        s.text(f'{c["pilotPass"]}/{c["total"]}',496,yy,58,9,'muted')
    note(s,'Barras con origen en cero y escala 0–100 %. Cada segmento tiene un denominador distinto.');sc.append(s)
    s=base(b,'report',5,'Medir también\nla experiencia de uso.','Latencia y progreso de iteraciones / datos ilustrativos.')
    s.text('Latencia por respuesta (segundos)',42,276,510,16,'ink','Medium')
    latency=e['latencySeconds']
    for i,(label,a,c) in enumerate([('p50',latency['baselineP50'],latency['pilotP50']),('p95',latency['baselineP95'],latency['pilotP95'])]):
        yy=322+i*83;s.text(label,42,yy,65,11,'ink','Medium');xx=113;ww=368
        s.box(xx,yy,ww*a/5,14,'muted');s.text(str(a).replace('.',','),493,yy,50,9,'muted')
        s.box(xx,yy+25,ww*c/5,14,'accent');s.text(str(c).replace('.',','),493,yy+25,50,9,'accent')
    s.text('Base / gris   ·   Piloto / berenjena   ·   escala 0–5 s',42,493,511,9,'muted')
    s.text('Calidad durante las iteraciones',42,542,511,16,'ink','Medium')
    x,y,w,h=55,602,484,90
    for value in (70,80,90):
        yy=y+h-(value-70)/25*h;s.line(x,yy,x+w,yy);s.text(str(value)+' %',42,yy-13,50,7,'muted')
    values=e['dailyQualityPercent']
    for i in range(len(values)-1):
        x1=x+i*w/6;y1=y+h-(values[i]-70)/25*h;x2=x+(i+1)*w/6;y2=y+h-(values[i+1]-70)/25*h
        s.line(x1,y1,x2,y2,'accent',1.8);s.circle(x1,y1,2.5,'accent')
    s.circle(x+w,y+h-(values[-1]-70)/25*h,2.5,'accent')
    for i in range(7):s.text(f'I{i+1}',x+i*w/6-2,y+h+13,27,8,'muted')
    s.text('Misma muestra en cada iteración; esta curva no mide generalización.',42,737,510,9,'muted')
    sc.append(s)
    s=base(b,'report',6,'Catorce fallos.\nCuatro causas accionables.','Clasificación exclusiva de los casos que no pasan / ejemplo sintético.')
    names=[('Documentación ausente','documentMissing'),('Respuesta incompleta','incompleteAnswer'),('Cita incorrecta','incorrectCitation'),('Respuesta fuera de alcance','unnecessaryAnswer')]
    for i,(label,key) in enumerate(names):
        yy=286+i*84;value=e['pilotErrorCounts'][key];s.text(label,42,yy,315,11,'ink','Medium');s.box(42,yy+25,455,16,'surface');s.box(42,yy+25,455*value/6,16,'accent');s.text(f'{value} casos',505,yy+24,49,9,'muted')
    s.box(42,656,511,92,'surface');s.text('Primera acción: completar las fuentes.',55,669,484,17,'accent','Medium');s.text('Seis fallos dependen de contenido ausente. Revisar las fuentes antes de ajustar prompts o añadir complejidad al sistema.',55,710,480,10,'muted')
    sc.append(s)
    s=base(b,'report',7,'Una decisión acompañada\nde un plan concreto.','Recomendaciones, responsables y condiciones de salida.')
    table(s,['Acción','Responsable','Validación'],[
      ['Completar las fuentes','Producto / cliente','Volver a evaluar los 6 casos de documentación.'],
      ['Ajustar recuperación y citas','Equipo técnico','Revisar casos complejos y trazas.'],
      ['Repetir evaluación','TenBeltz + cliente','Misma rúbrica; muestra ampliada y documentada.'],
      ['Abrir piloto controlado','Responsable de producto','Calidad acordada, observabilidad y retorno humano.']],42,278,[183,125,203],83)
    s.text('La entrega incluye la evidencia.',42,676,510,20,'accent','Medium');s.text('Conservar resultados por caso, versión de fuentes, configuración del sistema y límites. Acordar la decisión de lanzamiento con el cliente.',42,719,510,11,'muted',leading=1.45)
    sc.append(s)
    return sc


def docx_example(b,scenes,kind,d):
    doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=sec.bottom_margin=Mm(20);sec.left_margin=sec.right_margin=Mm(18)
    for name,size in [('Normal',10),('Title',30),('Heading 1',18)]:
        st=doc.styles[name];st.font.name=b.IDENTITY['font'];st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(b.C['ink'][1:])
    sec.header.paragraphs[0].add_run().add_picture(str(b.DIST/'logos'/'tenbeltz-horizontal-color.png'),width=Mm(44))
    sec.footer.paragraphs[0].text='TenBeltz / Ladera Cloud / ejemplo ficticio'
    for i,scene in enumerate(scenes):
        if i:doc.add_page_break()
        # Content stays editable. Complex diagrams are additionally reusable PNGs.
        resources={2:('architecture',279,638),3:('timeline',270,700)} if kind=='proposal' else {3:('quality',279,708),4:('latency',270,719),5:('errors',279,633)}
        resource=resources.get(i)
        title=True
        for node in scene.nodes:
            if node['type']!='text' or node['y']<100 or node['y']>778:continue
            if resource and resource[1]<=node['y']<=resource[2]:continue
            if node['size']>=25:
                doc.add_paragraph(node['text'],'Title' if title else 'Heading 1');title=False
            else:doc.add_paragraph(node['text'])
        if resource:
            doc.add_paragraph('Gráfico del ejemplo / datos ilustrativos:')
            doc.add_picture(str(b.DIST/'graphics'/f'tenbeltz-example-{resource[0]}.png'),width=Mm(155))
    doc.core_properties.author='TenBeltz';doc.core_properties.title='Ejemplo '+kind+' / Ladera Cloud'
    doc.save(b.DIST/'documents'/f'tenbeltz-{kind}-example.docx')


def build(b):
    d=json.loads((b.SOURCE/'example-project.json').read_text())
    directory=b.DIST/'documents'/'examples';directory.mkdir(exist_ok=True)
    for kind,fn in [('proposal',proposal),('report',report)]:
        pages=fn(b,d)
        target=b.DIST/'documents'/f'tenbeltz-{kind}-example.pdf'
        p=canvas.Canvas(str(target),pagesize=(595.276,841.89));p.setTitle('TenBeltz / '+kind+' / ejemplo ficticio');p.setAuthor('TenBeltz')
        for i,scene in enumerate(pages,1):
            scene.draw_pdf(p);p.showPage()
            b.export(f'tenbeltz-{kind}-example-{i:02d}',595.276,841.89,scene.svg_body(),f'{kind} / página {i}',('svg','png'),'documents/examples')
        p.save()
        # Export charts/diagrams separately for use in other branded documents.
        crops=[(3,279,359,'architecture'),(4,270,430,'timeline')] if kind=='proposal' else [(4,279,429,'quality'),(5,270,449,'latency'),(6,279,354,'errors')]
        for page_index,top,height,label in crops:
            body=f'<g transform="scale(4) translate(-42,-{top})">{pages[page_index-1].svg_body()}</g>'
            b.export(f'tenbeltz-example-{label}',2044,height*4,body,'Recurso del ejemplo ficticio / '+label,('svg','png'),'graphics')
        docx_example(b,pages,kind,d)
        # Readable offline presentation of all document pages, with explicit status.
        images=''.join(f'<article><img src="examples/tenbeltz-{kind}-example-{i:02d}.svg" alt="Página {i}: ejemplo ficticio de {kind}"></article>' for i in range(1,len(pages)+1))
        (b.DIST/'documents'/f'tenbeltz-{kind}-example.html').write_text(f'<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>TenBeltz / {kind} / ejemplo</title><style>body{{background:{b.C["surface"]};margin:0;font-family:Arial,sans-serif}}nav{{padding:24px;text-align:center;color:{b.C["ink"]}}}article{{width:min(100%,794px);margin:24px auto}}img{{width:100%}}a{{color:{b.C["accent"]}}}@media print{{nav{{display:none}}article{{margin:0;break-after:page}}}}@page{{size:A4;margin:0}}</style><body><nav>Ejemplo ficticio / datos ilustrativos · <a href="tenbeltz-{kind}-example.pdf">PDF</a> · <a href="tenbeltz-{kind}-example.docx">Word editable</a></nav>{images}</body></html>')
        # Contact sheet renders the actual native PDF, not an equivalent browser scene.
        images=[]
        with fitz.open(target) as pdf:
            for page in pdf:
                pix=page.get_pixmap(matrix=fitz.Matrix(.65,.65));im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);images.append(im)
        sheet=Image.new('RGB',(400*4,570*((len(images)+3)//4)),b.C['surface']);draw=ImageDraw.Draw(sheet)
        font=ImageFont.truetype(str(b.SOURCE/'fonts'/'IBMPlexSans-Regular.ttf'),14)
        for i,im in enumerate(images):sheet.paste(im,((i%4)*400,(i//4)*570));draw.text(((i%4)*400+8,(i//4)*570+547),f'{kind} / {i+1:02d}',font=font,fill=b.C['ink'])
        sheet.save(b.KIT/f'{kind}-example-contact-sheet.jpg',quality=93)
