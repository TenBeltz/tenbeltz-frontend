"""One geometry model shared by editable PPTX, SVG previews and native PDFs.

Coordinates are in design pixels; text is wrapped against real font metrics.
No slide is flattened to an image in PowerPoint.
"""
from pathlib import Path
import base64
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader


class Scene:
    def __init__(self,b,w=1600,h=900,bg='paper'):
        self.b=b;self.w=w;self.h=h;self.bg=bg;self.nodes=[]

    def box(self,x,y,w,h,fill='surface'):
        self.nodes.append(dict(type='rect',x=x,y=y,w=w,h=h,fill=fill));return self

    def circle(self,x,y,r,fill='surface'):
        self.nodes.append(dict(type='circle',x=x,y=y,r=r,fill=fill));return self

    def line(self,x,y,x2,y2,fill='line',stroke=1):
        self.nodes.append(dict(type='line',x=x,y=y,x2=x2,y2=y2,fill=fill,stroke=stroke));return self

    def text(self,s,x,y,w,size=28,fill='ink',weight='Regular',leading=1.3):
        lines=[]
        for p in str(s).split('\n'):
            cur=''
            for word in p.split():
                candidate=(cur+' '+word).strip()
                if cur and self.b.text_width(candidate,size,weight)>w:
                    lines.append(cur);cur=word
                else:cur=candidate
            lines.append(cur)
        height=len(lines)*size*leading
        if x<0 or y<0 or x+w>self.w+.1 or y+height>self.h+.1:
            raise ValueError(f'Text outside scene: {s!r}, {(x,y,w,height)}')
        self.nodes.append(dict(type='text',text='\n'.join(lines),x=x,y=y,w=w,h=height,size=size,fill=fill,weight=weight,leading=leading))
        return y+height

    def image(self,path,x,y,w,h):
        self.nodes.append(dict(type='image',path=str(path),x=x,y=y,w=w,h=h));return self

    def logo(self,x,y,w,variant='color'):
        lw,lh,_=self.b.logo_body('horizontal',variant)
        return self.image(self.b.DIST/'logos'/f'tenbeltz-horizontal-{variant}.png',x,y,w,w*lh/lw)

    def arrow(self,x,y,x2,y2,fill='accent',stroke=2):
        self.line(x,y,x2,y2,fill,stroke)
        if y==y2:self.line(x2-8,y2-5,x2,y2,fill,stroke);self.line(x2-8,y2+5,x2,y2,fill,stroke)
        else:self.line(x2-5,y2-8,x2,y2,fill,stroke);self.line(x2+5,y2-8,x2,y2,fill,stroke)

    def svg_body(self):
        b=self.b;out=b.rect(0,0,self.w,self.h,self.bg)
        for n in self.nodes:
            t=n['type']
            if t=='rect':out+=b.rect(n['x'],n['y'],n['w'],n['h'],n['fill'])
            elif t=='circle':out+=f'<circle cx="{n["x"]}" cy="{n["y"]}" r="{n["r"]}" fill="{b.color(n["fill"])}"/>'
            elif t=='line':out+=b.line(n['x'],n['y'],n['x2'],n['y2'],n['fill'],n['stroke'])
            elif t=='text':
                for i,line in enumerate(n['text'].split('\n')):
                    out+=b.text_svg(line,n['x'],n['y']+n['size']+i*n['size']*n['leading'],n['size'],n['fill'],n['weight'])
            elif t=='image':
                data=Path(n['path']).read_bytes();suffix=Path(n['path']).suffix[1:]
                mimetype='image/'+('svg+xml' if suffix=='svg' else suffix)
                out+=f'<image x="{n["x"]}" y="{n["y"]}" width="{n["w"]}" height="{n["h"]}" preserveAspectRatio="xMidYMid meet" href="data:{mimetype};base64,{base64.b64encode(data).decode()}"/>'
        return out

    def draw_pdf(self,p):
        b=self.b;H=self.h
        b.pdf_bg(p,self.w,H,self.bg)
        for n in self.nodes:
            t=n['type']
            if t=='rect':b.pdf_box(p,n['x'],n['y'],n['w'],n['h'],n['fill'],H)
            elif t=='circle':p.setFillColor(HexColor(b.color(n['fill'])));p.circle(n['x'],H-n['y'],n['r'],stroke=0,fill=1)
            elif t=='line':
                p.setStrokeColor(HexColor(b.color(n['fill'])));p.setLineWidth(n['stroke']);p.line(n['x'],H-n['y'],n['x2'],H-n['y2'])
            elif t=='text':
                for i,s in enumerate(n['text'].split('\n')):
                    b.pdf_text(p,s,n['x'],n['y']+n['size']+i*n['size']*n['leading'],n['size'],n['fill'],n['weight'],H)
            elif t=='image':b.pdf_image(p,n['path'],n['x'],n['y'],n['w'],n['h'],H)

    def draw_pptx(self,sl):
        # 120 design pixels/inch gives a standard 16:9 slide, 1600×900.
        px=lambda value:Inches(value/120)
        sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(self.b.color(self.bg)[1:])
        for n in self.nodes:
            t=n['type'];sh=None
            if t in ('rect','circle'):
                if t=='rect':x,y,w,h=n['x'],n['y'],n['w'],n['h'];kind=MSO_SHAPE.RECTANGLE
                else:x,y,w,h=n['x']-n['r'],n['y']-n['r'],n['r']*2,n['r']*2;kind=MSO_SHAPE.OVAL
                sh=sl.shapes.add_shape(kind,px(x),px(y),px(w),px(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(self.b.color(n['fill'])[1:]);sh.line.fill.background()
            elif t=='line':
                sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,px(n['x']),px(n['y']),px(n['x2']),px(n['y2']))
                sh.line.color.rgb=RGBColor.from_string(self.b.color(n['fill'])[1:]);sh.line.width=Pt(n['stroke']*.6)
            elif t=='image':sh=sl.shapes.add_picture(n['path'],px(n['x']),px(n['y']),width=px(n['w']),height=px(n['h']))
            elif t=='text':
                sh=sl.shapes.add_textbox(px(n['x']),px(n['y']),px(n['w']),px(n['h']+8))
                tf=sh.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                family=self.b.FONTS[n['weight']]['name'].getDebugName(1)
                for i,s in enumerate(n['text'].split('\n')):
                    para=tf.paragraphs[0] if i==0 else tf.add_paragraph();para.text=s
                    para.font.name=family;para.font.size=Pt(n['size']*.6);para.font.color.rgb=RGBColor.from_string(self.b.color(n['fill'])[1:]);para.space_after=Pt(0);para.line_spacing=Pt(n['size']*.6*n['leading'])
            if sh is not None:sh.name=f'TenBeltz / {t} / {len(sl.shapes):02d}'


def slide(b,d,i,n):
    kind=d['kind'];dark=kind in ('workshop','close','section');bg='accent' if kind=='workshop' else ('ink' if dark else 'paper')
    fg='white' if dark else 'ink';secondary='petal' if dark else 'muted'
    s=Scene(b,bg=bg);s.logo(78,56,235,'white' if dark else 'color');s.text(b.IDENTITY['website'],1290,67,235,18,secondary)
    s.line(78,817,1522,817,'muted' if dark else 'line');s.text('TENBELTZ / BRAND SYSTEM 2026.10',78,842,950,15,secondary);s.text(f'{i:02d} / {n:02d}',1420,842,95,15,secondary)
    blocks=d.get('blocks',[])
    # Intentional rhythm: large type, structural diagrams, light/dark sections.
    if kind=='cover':
        s.image(b.DIST/'graphics'/'tenbeltz-flower-transparent.png',925,125,640,640)
        s.text('INGENIERÍA DE IA / AI ENGINEERING',78,203,700,18,'muted','Medium')
        s.text(d['title'],78,294,800,80,'ink','Medium',1.08)
        s.text(d.get('subtitle',''),78,562,755,28,'muted')
    elif kind=='statement':
        s.text(d['title'],78,233,705,72,'ink','Medium',1.1)
        s.text(d.get('subtitle',''),78,554,650,28,'muted')
        s.box(880,191,642,539,'surface')
        for j,(t,p) in enumerate(blocks):
            yy=238+j*235;s.text(t,930,yy,540,20,'accent','Medium');s.text(p,930,yy+51,530,30)
    elif kind in ('pathways','three'):
        s.text(d['title'],78,190,1390,55,'ink','Medium')
        for j,(t,p) in enumerate(blocks):
            xx=78+j*490;s.text(f'0{j+1}',xx,326,390,100,'petal','Medium');s.line(xx,458,xx+430,458)
            s.text(t,xx,488,420,46,'accent','Medium');s.text(p,xx,579,405,28,'muted')
    elif kind=='blueprint':
        s.text(d['title'],78,181,1450,54,'ink','Medium')
        for j,(t,p) in enumerate(blocks):
            xx=78+j*750;s.text(t,xx,282,670,30,'accent','Medium');s.text(p,xx,335,660,26,'muted')
        nodes=d.get('diagramLabels',['Datos','Recuperación','Sistema de IA','Evaluación'])
        for j,t in enumerate(nodes):
            xx=78+j*373;s.box(xx,550,324,116,'surface');s.text(t,xx+22,582,282,29,'ink','Medium')
            if j<3:s.arrow(xx+334,608,xx+362,608)
        s.text(d.get('caption','Flujo simplificado / arquitectura conceptual'),78,704,1400,18,'muted')
    elif kind=='delivery':
        s.text(d['title'],78,190,1290,57,'ink','Medium')
        s.image(b.DIST/'graphics'/'tenbeltz-background-routed-lines-paper.png',910,316,610,390)
        for j,(t,p) in enumerate(blocks):
            yy=330+j*205;s.text(f'0{j+1}',78,yy,95,50,'accent','Medium');s.text(t,204,yy,650,34,'ink','Medium');s.text(p,204,yy+63,640,27,'muted')
    elif kind=='workshop':
        s.text(d['title'],78,199,930,75,'white','Medium',1.1)
        s.text(d.get('subtitle',''),78,450,790,31,'petal')
        s.image(b.DIST/'graphics'/'tenbeltz-motif-line-petals-white.png',990,170,520,520)
        s.line(78,662,885,662,'petal')
        s.text(d.get('caption','Aprender / Probar / Aplicar'),78,695,1100,30,'white','Medium')
    elif kind in ('timeline','process'):
        s.text(d['title'],78,190,1440,58,'ink','Medium');s.line(172,412,1430,412,'accent',2)
        for j,(t,p) in enumerate(blocks):
            xx=78+j*375;s.circle(xx+20,412,10,'accent');s.text(f'0{j+1}',xx,313,280,56,'accent','Medium')
            s.text(t.split(' / ')[-1],xx,463,313,31,'ink','Medium');s.text(p,xx,547,300,27,'muted')
    elif kind=='evaluation':
        s.text(d['title'],78,190,1370,57,'ink','Medium')
        for j,(t,p) in enumerate(blocks):
            xx=115+j*492;s.circle(xx+155,466,143,'surface' if j!=1 else 'petal');s.text(t,xx+23,434,290,34,'accent','Medium')
            s.text(p,xx-10,654,365,26,'muted')
    elif kind=='responsibility':
        s.text(d['title'],78,190,1390,58,'ink','Medium')
        s.box(78,335,1444,360,'surface');s.line(741,374,741,650)
        for j,(t,p) in enumerate(blocks):
            xx=122+j*740;s.text(t,xx,380,580,25,'accent','Medium');s.text(p,xx,438,580,34,'ink','Medium')
        s.text(d.get('caption','Una dirección técnica / un equipo de implementación'),78,742,1390,20,'muted')
    elif kind=='evidence':
        s.text(d['title'],78,190,1410,62,'ink','Medium')
        for j,t in enumerate(d.get('labels',['Contexto','Contribución','Evidencia'])):
            yy=345+j*130;s.text(f'0{j+1}',78,yy,100,25,'accent','Medium');s.text(t,212,yy-9,450,44,'ink','Medium');s.line(78,yy+82,702,yy+82)
        for j,(t,p) in enumerate(blocks):
            yy=343+j*206;s.text(t,885,yy,625,23,'accent','Medium');s.text(p,885,yy+51,600,29,'muted')
    elif kind=='conversation':
        s.text(d['title'],78,190,1430,57,'ink','Medium')
        for j,(t,p) in enumerate(blocks):
            yy=337+j*131;s.text(t,78,yy,430,34,'accent','Medium');s.text(p,600,yy+3,880,30,'muted');s.line(78,yy+97,1522,yy+97)
    elif kind in ('section','close'):
        s.image(b.DIST/'graphics'/'tenbeltz-motif-line-petals-white.png',1065,265,425,425)
        s.text(d['title'],78,239,1110,78,'white','Medium',1.1)
        s.line(78,533,930,533,'muted');s.text(d.get('subtitle',''),78,586,960,36,'petal')
    elif kind=='table':
        s.text(d['title'],78,190,1450,54,'ink','Medium')
        for row,values in enumerate(d['rows']):
            yy=341+row*95
            for j,t in enumerate(values):
                xx=78+j*481;s.box(xx,yy,479,90,'accent' if row==0 else 'surface');s.text(t,xx+24,yy+24,430,28,'white' if row==0 else 'ink')
    elif kind=='metric':
        s.text(d['title'],78,190,1400,55,'ink','Medium');s.text(d.get('value','[Valor]'),78,350,950,156,'accent','Medium');s.text(d.get('subtitle',''),78,620,1240,29,'muted')
    elif kind=='image':
        s.text(d['title'],78,180,1430,55,'ink','Medium');s.box(78,312,1444,405,'surface');s.text(d.get('subtitle',''),118,400,1240,35,'muted')
    else: # retained comparison layout for the editable case template only
        s.text(d['title'],78,190,1400,57,'ink','Medium')
        for j,(t,p) in enumerate(blocks):
            xx=78+j*748;s.text(t,xx,359,660,34,'accent','Medium');s.text(p,xx,435,670,31,'muted')
    return s


def make_pptx(b,data,name):
    prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
    for rel in prs.slide_masters[0].part.rels.values():
        if rel.reltype.endswith('/theme'):
            root=ET.fromstring(rel.target_part.blob)
            for e in root.iter():
                if e.tag.endswith('latin'):e.set('typeface',b.IDENTITY['font'])
            rel.target_part._blob=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    for i,d in enumerate(data,1):
        sl=prs.slides.add_slide(prs.slide_layouts[6]);slide(b,d,i,len(data)).draw_pptx(sl)
        sl.notes_slide.notes_text_frame.text='TenBeltz / composición '+d['kind']+'. Texto, formas y conectores editables. Datos compartidos con SVG/HTML/PDF en source/slides.json. Instalar IBM Plex Sans antes de editar.'
    prs.core_properties.title='TenBeltz / '+name;prs.core_properties.author='TenBeltz'
    prs.save(b.DIST/'presentations'/f'tenbeltz-{name}.pptx')


def make_pdf(b,data,name):
    p=canvas.Canvas(str(b.DIST/'presentations'/f'tenbeltz-{name}.pdf'),pagesize=(960,540))
    p.setTitle('TenBeltz / '+name);p.setAuthor('TenBeltz')
    for i,d in enumerate(data,1):
        p.saveState();p.scale(.6,.6);slide(b,d,i,len(data)).draw_pdf(p);p.restoreState();p.showPage()
    p.save()
